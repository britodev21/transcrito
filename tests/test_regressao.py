"""
Regressao dos exemplos que funcionam: o que sai hoje tem que bater com saidas/.

Os padroes de valor e de codigo sao usados pelos dois extratores, entao mexer
no do holerite pode quebrar o cartao de ponto e vice-versa. saidas/ e a
referencia conferida a mao contra o PDF; qualquer diferenca aqui e mudanca de
comportamento, e precisa ser de proposito.

Quando a mudanca for de proposito, o gerar_saidas.py refaz saidas/ e o diff da
planilha entra no mesmo commit que a mudanca no extrator.
"""

import pytest
from openpyxl import load_workbook

from backend.extrator import processar_cartao_ponto
from backend.extrator_holerite import processar_holerite
from backend.planilha import GERADORES
from tests.casos import DIRETORIO_EXEMPLOS, DIRETORIO_SAIDAS, ler_saida

EXTRATORES = {
    "cartao-ponto": processar_cartao_ponto,
    "holerite": processar_holerite,
}

# Texto nativo: rodam em qualquer maquina, sem tesseract.
NATIVOS = [
    ("time-card-01", "cartao-ponto"),
    ("payroll-02", "holerite"),
    ("payroll-03", "holerite"),
]

# Escaneados: dependem do tesseract com portugues e levam minutos.
ESCANEADOS = [
    ("time-card-02", "cartao-ponto"),
]


def tesseract_com_portugues():
    try:
        import pytesseract

        return "por" in pytesseract.get_languages()
    except Exception:
        return False


@pytest.fixture(scope="module")
def extraidos():
    """Cache por modulo: extrair e a parte lenta, e tres testes usam o mesmo."""
    cache = {}

    def extrair(nome, tipo):
        if nome not in cache:
            cache[nome] = EXTRATORES[tipo](str(DIRETORIO_EXEMPLOS / f"{nome}.pdf"))
        return cache[nome]

    return extrair


def celulas_e_cores(caminho):
    """
    O xlsx salvo nao e comparavel byte a byte: o arquivo leva data de criacao.
    Compara o que o usuario ve, valor e preenchimento de cada celula.
    """
    aba = load_workbook(caminho).active
    return [
        [(celula.value, celula.fill.start_color.rgb) for celula in linha]
        for linha in aba.iter_rows()
    ]


@pytest.mark.parametrize(("nome", "tipo"), NATIVOS, ids=[n for n, _t in NATIVOS])
def test_json_igual_ao_de_saidas(nome, tipo, extraidos):
    assert extraidos(nome, tipo) == ler_saida(nome)


@pytest.mark.parametrize(("nome", "tipo"), NATIVOS, ids=[n for n, _t in NATIVOS])
def test_csv_igual_ao_de_saidas(nome, tipo, extraidos, tmp_path):
    caminho = tmp_path / f"{nome}.csv"
    GERADORES[tipo]["csv"](extraidos(nome, tipo), str(caminho))

    assert caminho.read_bytes() == (DIRETORIO_SAIDAS / f"{nome}.csv").read_bytes()


@pytest.mark.parametrize(("nome", "tipo"), NATIVOS, ids=[n for n, _t in NATIVOS])
def test_xlsx_igual_ao_de_saidas(nome, tipo, extraidos, tmp_path):
    caminho = tmp_path / f"{nome}.xlsx"
    GERADORES[tipo]["xlsx"](extraidos(nome, tipo), str(caminho))

    assert celulas_e_cores(caminho) == celulas_e_cores(
        DIRETORIO_SAIDAS / f"{nome}.xlsx"
    )


@pytest.mark.ocr
@pytest.mark.skipif(
    not tesseract_com_portugues(),
    reason="tesseract com o idioma portugues nao instalado",
)
@pytest.mark.parametrize(
    ("nome", "tipo"), ESCANEADOS, ids=[n for n, _t in ESCANEADOS]
)
def test_escaneado_igual_ao_de_saidas(nome, tipo, extraidos):
    """
    Fora da rodada padrao (pytest -m ocr pra rodar). Alem de lento, o OCR
    depende da versao do tesseract: saidas/ foi gerado com a 5.4.0, e outra
    versao pode ler um digito diferente sem que o extrator tenha mudado.
    """
    assert extraidos(nome, tipo) == ler_saida(nome)
