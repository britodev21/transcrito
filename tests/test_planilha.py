"""
Regras de destaque conferidas no arquivo salvo, e nao no codigo que as aplica.

Ler a cor de volta do xlsx e o que prova que o destaque chega ao usuario: um
aviso calculado certo mas pintado na celula errada passaria num teste que so
olhasse o dicionario de avisos.
"""

import csv

import pytest
from openpyxl import load_workbook

from backend.planilha import (
    COR_AMARELO,
    COR_BORDA_VERMELHA,
    COR_CABECALHO,
    COR_VERMELHO,
    GERADORES,
)
from tests.casos import (
    AMARELO,
    CASOS_CARTAO_PONTO,
    CASOS_HOLERITE,
    VERMELHO,
    todos_os_casos,
)

CASOS = list(todos_os_casos())


def cor(celula):
    """O openpyxl devolve ARGB ("00FFF3CD"); a planilha.py guarda so o RGB."""
    if celula.fill.fill_type is None:
        return None
    return celula.fill.start_color.rgb[-6:]


def destaque_da_linha(aba, numero_linha):
    """
    Le o destaque de uma linha e confere que ele cobre a largura toda.

    Linha meio pintada falha aqui: na planilha do holerite a maioria das
    colunas fica vazia, e e justamente nelas que o preenchimento costuma faltar.
    """
    cores = {
        cor(aba.cell(row=numero_linha, column=coluna))
        for coluna in range(1, aba.max_column + 1)
    }
    assert len(cores) == 1, f"linha {numero_linha} pintada pela metade: {cores}"

    return {COR_VERMELHO: VERMELHO, COR_AMARELO: AMARELO, None: None}[cores.pop()]


def gerar(tipo, formato, value, tmp_path):
    caminho = tmp_path / f"saida.{formato}"
    GERADORES[tipo][formato](value, str(caminho))
    return caminho


@pytest.mark.parametrize(
    ("tipo", "value", "esperado"),
    [(tipo, value, esperado) for _nome, tipo, value, esperado in CASOS],
    ids=[nome for nome, *_resto in CASOS],
)
def test_destaques_lidos_do_xlsx(tipo, value, esperado, tmp_path):
    aba = load_workbook(gerar(tipo, "xlsx", value, tmp_path)).active

    lidos = [destaque_da_linha(aba, linha) for linha in range(2, aba.max_row + 1)]
    assert lidos == esperado

    # A borda vermelha na primeira coluna acompanha so a linha vermelha.
    for numero_linha, destaque in enumerate(lidos, start=2):
        borda = aba.cell(row=numero_linha, column=1).border.left
        if destaque == VERMELHO:
            assert borda.style == "medium"
            assert borda.color.rgb[-6:] == COR_BORDA_VERMELHA
        else:
            assert borda.style is None


@pytest.mark.parametrize("tipo", ["cartao-ponto", "holerite"])
def test_cabecalho_na_cor_da_marca(tipo, tmp_path):
    _nome, _tipo, value, _esperado = next(c for c in CASOS if c[1] == tipo)
    aba = load_workbook(gerar(tipo, "xlsx", value, tmp_path)).active

    for celula in aba[1]:
        assert cor(celula) == COR_CABECALHO
        assert celula.font.bold
        assert celula.font.color.rgb[-6:] == "FFFFFF"


def test_colunas_do_cartao_ponto_seguem_o_dia_com_mais_batidas(tmp_path):
    value, _esperado = CASOS_CARTAO_PONTO["batida_impar"]
    aba = load_workbook(gerar("cartao-ponto", "xlsx", value, tmp_path)).active

    # Tres batidas pedem dois pares: a Saida 2 fica vazia, mas existe.
    assert [c.value for c in aba[1]] == [
        "Data", "Entrada 1", "Saída 1", "Entrada 2", "Saída 2",
    ]


def test_holerite_com_folha_ganha_coluna_e_congela_depois_dela(tmp_path):
    value, _esperado = CASOS_HOLERITE["duas_folhas_do_mesmo_mes"]
    aba = load_workbook(gerar("holerite", "xlsx", value, tmp_path)).active

    assert [c.value for c in aba[1]][:4] == ["Pág.", "Folha", "Mês", "Ano"]
    assert aba.freeze_panes == "E2"


def test_holerite_sem_folha_nao_tem_coluna_folha(tmp_path):
    value, _esperado = CASOS_HOLERITE["mes_pulado"]
    aba = load_workbook(gerar("holerite", "xlsx", value, tmp_path)).active

    assert [c.value for c in aba[1]][:3] == ["Pág.", "Mês", "Ano"]
    assert aba.freeze_panes == "D2"


def test_verba_repetida_ganha_contador_em_vez_de_sobrescrever(tmp_path):
    value, _esperado = CASOS_HOLERITE["verba_repetida_na_mesma_folha"]
    aba = load_workbook(gerar("holerite", "xlsx", value, tmp_path)).active

    linha = dict(zip([c.value for c in aba[1]], [c.value for c in aba[2]], strict=True))
    assert linha["CONTRIBUICAO NEGOCIAL"] == "-10,00"
    assert linha["CONTRIBUICAO NEGOCIAL (2)"] == "-12,00"


def test_mes_e_ano_continuam_texto(tmp_path):
    """Virar numero comeria o zero a esquerda: "01" sairia como 1."""
    value, _esperado = CASOS_HOLERITE["mes_pulado"]
    aba = load_workbook(gerar("holerite", "xlsx", value, tmp_path)).active

    assert aba["B2"].value == "01"
    assert aba["C2"].value == "2018"


@pytest.mark.parametrize(
    ("tipo", "value"),
    [(tipo, value) for _nome, tipo, value, _esperado in CASOS],
    ids=[nome for nome, *_resto in CASOS],
)
def test_csv_tem_as_mesmas_celulas_do_xlsx(tipo, value, tmp_path):
    aba = load_workbook(gerar(tipo, "xlsx", value, tmp_path)).active
    do_xlsx = [
        ["" if celula.value is None else str(celula.value) for celula in linha]
        for linha in aba.iter_rows()
    ]

    caminho_csv = gerar(tipo, "csv", value, tmp_path)

    # BOM pro Excel acertar a acentuacao, ";" pro Excel em portugues separar.
    assert caminho_csv.read_bytes().startswith(b"\xef\xbb\xbf")
    with open(caminho_csv, encoding="utf-8-sig", newline="") as arquivo:
        do_csv = list(csv.reader(arquivo, delimiter=";"))

    assert do_csv == do_xlsx
